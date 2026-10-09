"""Schema 2: atomic customer view bound to request and exact accepted authority."""
from __future__ import annotations
import csv, json, os, uuid
from collections import defaultdict
from datetime import timedelta
from pathlib import Path
from typing import Literal
from shelfcash_forecast.optimization.scenario_contracts import ScenarioContract
from shelfcash_forecast.optimization.planning_service import content_hash
from shelfcash_forecast.optimization.shipment_fees import shipment_group
from shelfcash_forecast.optimization.export import ORDER_FIELDS
from shelfcash_forecast.json_output import write_json

LEDGER_FIELDS=['scenario_id','probability_weight','store_id','ingredient_id','ingredient_name','unit','date',
    'beginning_quantity','inbound_quantity','maximum_quantity','usable_before_consumption',
    'demand_quantity','fulfilled_quantity','shortage_quantity','expired_quantity','waste_quantity','ending_quantity',
    'at_risk_expiry_quantity','capacity_violation_quantity','holding_cost','shortage_cost','expiry_cost','waste_cost',
    'receiving_occupancy_liters','ending_occupancy_liters','lots_consumed','lots_expired']
CUSTOMER_ORDER_FIELDS=ORDER_FIELDS+['store_name','shipment_group_id','pack_unit','classification','contract_id']

class CustomerProcurementPackage(ScenarioContract):
    schema_version: Literal[2]=2
    request_id: str
    plan_id: str | None
    profile_id: str | None
    profile_version: int | None
    planning_mode: str
    environment: str
    data_origin: str
    technical_outcome: str
    critic_passed: bool
    business_ready: Literal[False]=False
    execution_authorized: Literal[False]=False
    procurement_executed: Literal[False]=False
    operational_status: Literal['NOT_FOR_OPERATION']='NOT_FOR_OPERATION'
    metadata: dict
    accepted_plan: dict | None
    orders: list[dict]
    assumptions_and_customer_conditions: list[dict]
    cost_decomposition: dict
    service_risk_metrics: dict | None
    rule_coverage: list[dict]
    accepted_zero_purchase: bool
    diagnostics: list[dict]

def display_name(request,kind,identifier,store=None,fallback=None):
    registry=request.entity_display_names.get(kind,{})
    return registry.get(f'{store}|{identifier}',registry.get(identifier,fallback or identifier))

def csv_rows(path,rows,fields=None):
    fields=fields or (list(rows[0]) if rows else [])
    if not fields: raise ValueError('CSV_SCHEMA_REQUIRED')
    with Path(path).open('w',encoding='utf-8-sig',newline='') as handle:
        writer=csv.DictWriter(handle,fieldnames=fields,extrasaction='ignore');writer.writeheader()
        for row in rows:
            values={k:json.dumps(v,ensure_ascii=False,default=str) if isinstance(v,(dict,list)) else v for k,v in row.items()}
            values={k:"'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for k,v in values.items()}
            writer.writerow(values)

def perishable_evidence(request,candidate):
    rows=[];terminal=[];weighted=defaultdict(float);worst=defaultdict(float)
    if candidate is None or candidate.simulation is None:
        return {'status':'UNAVAILABLE_NO_ACCEPTED_PLAN','terminal_lots':[],'by_key':[]},rows
    for world in candidate.simulation.results:
        for ledger in world.daily_ledgers:
            d=ledger.model_dump(mode='json');d['date']=d.pop('simulation_date');key=(ledger.store_id,ledger.ingredient_id)
            coefficients=[c.liters_per_base_unit for r in request.normalized_rules
                if r.store_id==key[0] and r.semantics=='GLOBAL_RECEIVING_PEAK' and r.active(ledger.simulation_date)
                for c in r.occupancy if c.ingredient_id==key[1] and c.base_unit==ledger.unit]
            coefficient=coefficients[0] if coefficients else None
            traces=[t.model_dump(mode='json') for t in world.consumption_traces if (t.store_id,t.ingredient_id,t.simulation_date)==(*key,ledger.simulation_date)]
            expired=[t.model_dump(mode='json') for t in world.expiry_traces if (t.store_id,t.ingredient_id,t.simulation_date)==(*key,ledger.simulation_date)]
            balance=ledger.beginning_quantity+ledger.inbound_quantity-ledger.fulfilled_quantity-ledger.expired_quantity-ledger.waste_quantity
            if abs(balance-ledger.ending_quantity)>request.inventory_policy.accounting_tolerance: raise ValueError('CUSTOMER_LEDGER_ACCOUNTING_MISMATCH')
            d.update(probability_weight=world.probability_weight,ingredient_name=display_name(request,'ingredients',key[1],key[0]),
                usable_before_consumption=ledger.beginning_quantity+ledger.inbound_quantity-ledger.expired_quantity,
                receiving_occupancy_liters=None if coefficient is None else coefficient*ledger.maximum_quantity,
                ending_occupancy_liters=None if coefficient is None else coefficient*ledger.ending_quantity,
                lots_consumed=traces,lots_expired=expired)
            rows.append(d)
        per_world=defaultdict(float)
        for lot in world.ending_lots:
            if lot.expiry_date is None: continue
            remaining_days=(lot.expiry_date-request.planning_end_date).days
            item={**lot.model_dump(mode='json'),'scenario_id':world.scenario_id,'probability_weight':world.probability_weight,
                'ingredient_name':display_name(request,'ingredients',lot.ingredient_id,lot.store_id),
                'expiry_days_after_horizon':remaining_days,'at_risk_terminal':remaining_days<=request.inventory_policy.at_risk_expiry_days,
                'certain_post_horizon_waste':False}
            terminal.append(item)
            if item['at_risk_terminal']: per_world[(lot.store_id,lot.ingredient_id,lot.unit)]+=lot.quantity_remaining
        for key,q in per_world.items():
            if world.probability_weight is not None: weighted[key]+=world.probability_weight*q
            worst[key]=max(worst[key],q)
    return {'status':'EXACT_SIMULATED','expiry_inclusive':request.inventory_policy.expiry_inclusive,
        'shelf_life_convention':'arrival + calendar-day offset; inclusive/exclusive from request policy',
        'post_horizon_simulation':False,'terminal_stock_is_not_certain_waste':True,'at_risk_window_days':request.inventory_policy.at_risk_expiry_days,
        'terminal_lots':terminal,'by_key':[{'store_id':k[0],'ingredient_id':k[1],'unit':k[2],
            'expected_at_risk_terminal':weighted[k] if all(w.probability_weight is not None for w in candidate.simulation.results) else None,
            'worst_at_risk_terminal':worst[k]} for k in sorted(worst)]},rows

def build_customer_view(run,accepted):
    request,result=run.request,run.result
    candidate=result.evaluations.get(result.recommended_strategy) if result.recommended_strategy else None
    if candidate and (not candidate.critic.passed or not candidate.plan.completed or not candidate.critic.checks.get('evaluation_coverage')):
        raise ValueError('CUSTOMER_EXPORT_REQUIRES_FULL_EXACT_ACCEPTANCE')
    offers={o.offer_id:o for o in request.supplier_offers};rows=[]
    for row in accepted['orders']:
        offer=offers[row['opportunity_id']];terms=offer.source_terms
        rows.append({**row,'ingredient_name':display_name(request,'ingredients',offer.ingredient_id,offer.store_id,terms.get('ingredient_name')),
            'supplier_name':display_name(request,'suppliers',offer.supplier_id,offer.store_id,terms.get('supplier_name')),
            'store_name':display_name(request,'stores',offer.store_id),'shipment_group_id':shipment_group(offer),
            'pack_unit':offer.unit,'classification':terms.get('classification','REQUEST_DECLARED_TERMS'),'contract_id':terms.get('contract_id')})
    purchase=sum(r['purchase_cost'] for r in rows);delivery=sum(r['allocated_delivery_cost'] for r in rows)
    full=request.evaluation_scenarios or request.demand_scenarios;binding=request.planning_binding
    simulation=candidate.simulation if candidate else None
    metrics=simulation.risk_metrics.model_dump(mode='json') if simulation and simulation.risk_metrics else None
    rules=candidate.critic.details.get('business_rule_evaluation',[]) if candidate else request.rule_coverage
    diagnostics=[d.model_dump(mode='json') for d in result.diagnostics]
    if not candidate:
        for strategy,e in result.evaluations.items():
            diagnostics.append({'reason_code':'CANDIDATE_REJECTED','strategy':strategy,'proof_status':'OBSERVED',
                'solver_status':e.plan.solver_status,'violations':e.critic.hard_violations,
                'scope':e.plan.provenance.get('optimization_scope'),'search_attempts':[{k:a.get(k) for k in ('phase','termination','optimization_ids')} for a in e.attempts]})
    package=CustomerProcurementPackage(request_id=request.request_id,plan_id=candidate.plan.plan_id if candidate else None,
        profile_id=binding.get('profile_id'),profile_version=binding.get('profile_version'),planning_mode=request.planning_mode,environment=request.environment,
        data_origin=run.provenance.get('data_origin','REQUEST_DECLARED_INPUT'),technical_outcome=result.technical_outcome,critic_passed=bool(candidate),
        metadata={'snapshot_date':str(request.inventory_snapshot_date) if request.inventory_snapshot_date else None,
            'decision_date':str(request.decision_date),'snapshot_boundary':request.inventory_snapshot_boundary,
            'planning_horizon':{'start':str(min(s.simulation_start_date or min(l.target_date for l in s.lines) for s in full)),'end':str(request.planning_end_date)},'timezone':request.timezone,
            'request_hash':content_hash(request),'config_hash':content_hash(run.config),'profile_binding':binding,'source_bundle':run.provenance.get('source_bundle'),
            'full_pool_hash':content_hash([s.model_dump(mode='json') for s in full]),'world_count':len(full),'world_weights':{s.scenario_id:s.probability_weight for s in full},
            'evaluation_coverage':len(simulation.results) if simulation else 0,'optimization_world_count':len(request.demand_scenarios),'selected_mode':run.selected_mode,
            'selected_strategy':result.recommended_strategy,'solver_status':candidate.plan.solver_status if candidate else None,
            'requested_mode':'stochastic' if request.stochastic else 'deterministic',
            'optimizer_actual_mode':result.provenance.get('actual_mode'),
            'solver_physics_mode':candidate.plan.provenance.get('mode') if candidate else None,
            'mip_gap':candidate.plan.provenance.get('mip_gap') if candidate else None,'optimization_scope':candidate.plan.provenance.get('optimization_scope') if candidate else None,
            'search_phase':candidate.plan.provenance.get('search_phase') if candidate else None,
            'budget':{'amount':request.budget,'cap_enabled':request.budget is not None,'currency':request.currency,'scope':request.budget_scope},
            'budget_satisfied':bool(candidate) and (request.budget is None or purchase+delivery<=request.budget+request.inventory_policy.accounting_tolerance),
            'shipment_count':len({r['shipment_group_id'] for r in rows}),
            'metric_definitions':{'mean_key_fill_rate':'mean of scenario-weighted key fill rates; zero-demand key fill=1','any_stockout_probability':'weight of worlds with shortage in any key; not OOS guarantee'},
            'what_if_lineage':{k:run.provenance.get(k) for k in ('baseline_path','what_if_id','demand_origin','config_role')},'scenario_provenance':request.scenario_provenance},
        accepted_plan=candidate.plan.model_dump(mode='json') if candidate else None,orders=rows,assumptions_and_customer_conditions=binding.get('assumptions',[]),
        cost_decomposition={'purchase_vnd':purchase,'delivery_vnd':delivery,'committed_procurement_vnd':purchase+delivery,'cost_currency':request.currency,
            'cost_components_by_optimization_world':candidate.plan.provenance.get('cost_components_by_world',{}) if candidate else {},
            'objective_value':candidate.plan.objective_value if candidate else None,'soft_rule_penalty_vnd':sum(r.get('soft_penalty_vnd',0) for r in rules)},
        service_risk_metrics=metrics,rule_coverage=rules,accepted_zero_purchase=bool(candidate and not rows),diagnostics=diagnostics)
    return package,candidate

def explanation(package):
    m=package.metadata;b=m['budget'];h=m['planning_horizon']
    lines=['# Kết quả lập kế hoạch ShelfCash',f"Trạng thái kỹ thuật: **{package.technical_outcome}**; critic: {'PASS' if package.critic_passed else 'chưa có plan được chấp nhận'}.",
        f"Request: `{package.request_id}`. Profile: {package.profile_id or 'typed request'} v{package.profile_version}. Snapshot: {m['snapshot_date']} ({m['snapshot_boundary']}); horizon {h['start']} – {h['end']}; timezone {m['timezone']}.",
        f"Budget: {'không có cap' if b['amount'] is None else str(b['amount'])}; currency {b['currency']}; scope `{b['scope']}`.",
        f"Pool: {m['world_count']} worlds; exact coverage: {m['evaluation_coverage']}. IDs/weights trong JSON; chưa xác lập độ chính xác ngoài mẫu.",'Business ready=false; execution authorized=false; chưa gửi đơn.']
    if package.critic_passed:
        lines.append('Không cần mua trong scope đã đánh giá.' if package.accepted_zero_purchase else 'Các orders dưới đây đã được exact critic chấp nhận.')
        lines.append(f"Mua + giao: {package.cost_decomposition['committed_procurement_vnd']:,.2f} {b['currency']}. Holding/shortage/expiry/soft costs và objective/search scope báo riêng trong JSON.")
        if package.service_risk_metrics:
            metrics=package.service_risk_metrics;lines.append(f"Mean key fill: {metrics['mean_key_fill_rate']:.6%}; any-stockout pool probability: {metrics['any_stockout_probability']:.6%}.")
        lines+=['| Ingredient / ID | Supplier / ID | Đặt | Nhận | Expiry | Quantity | Packs | Cash |','|---|---|---|---|---|---:|---:|---:|']
        def esc(s):return str(s).replace('|','\\|').replace('\n',' ')
        for r in package.orders:lines.append(f"| {esc(r['ingredient_name'])} / {esc(r['ingredient_id'])} | {esc(r['supplier_name'])} / {esc(r['supplier_id'])} | {r['order_date']} | {r['arrival_date']} | {r['expiry_date']} | {r['quantity']:g} {r['unit']} | {r['pack_count']} | {r['total_procurement_cost']:,.2f} |")
    else:
        lines.append('Không có accepted orders. Orders rỗng ở đây không phải chứng nhận zero-purchase. Candidate FAIL không được dùng làm đơn mua.')
        lines += [f"- {d['reason_code']}: {d.get('field_paths',[])}; {d.get('expected_meaning') or d.get('action_required') or d.get('violations')}" for d in package.diagnostics]
    lines+=['Daily Ledger và terminal_expiry.json phân biệt expiry trong horizon với tồn cuối có rủi ro; không mô phỏng sau horizon.',
        'Assumptions/terms cần xác nhận: Conditions và plan_conditions.json. Fixed commitments hoặc feasibility phase không chứng minh global cost optimum.']
    return '\n\n'.join(lines)+'\n'

def write_workbook(path,package,rows,conditions,ledger,terminal):
    from datetime import date,datetime
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook();wb.remove(wb.active);expected={}
    def sheet(name,headers,data):
        ws=wb.create_sheet(name);ws.append(headers);expected[name]=[headers]+data
        for row in data:ws.append([json.dumps(v,ensure_ascii=False,default=str) if isinstance(v,(dict,list)) else v for v in row])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F4E78');c.alignment=Alignment(wrap_text=True)
        for j,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(j)].width=min(55,max(18,len(str(h))+3))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                h=str(headers[cell.column-1]);cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,str):cell.data_type='s'
                if h in {'date','order_date','arrival_date','expiry_date'} and isinstance(cell.value,str):
                    cell.value=date.fromisoformat(cell.value);cell.number_format='yyyy-mm-dd'
                if isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool):cell.number_format='0.0000%' if any(t in h for t in ('fill_rate','probability')) else '#,##0.000000'
            import math
            estimated_lines=max((math.ceil(len(str(c.value or ''))/max(10,ws.column_dimensions[get_column_letter(c.column)].width-2)) for c in row),default=1)
            ws.row_dimensions[row[0].row].height=min(409,max(30,15*(estimated_lines+1)))
    m=package.metadata
    sheet('Summary',['Field','Value'],[[k,v] for k,v in {'Status':package.technical_outcome,'Request':package.request_id,'Profile':package.profile_id,'Profile version':package.profile_version,
        'Snapshot':m['snapshot_date'],'Snapshot boundary':m['snapshot_boundary'],'Horizon start':m['planning_horizon']['start'],'Horizon end':m['planning_horizon']['end'],'Timezone':m['timezone'],
        'World count':m['world_count'],'Exact coverage':m['evaluation_coverage'],'Critic PASS':package.critic_passed,'Accepted zero purchase':package.accepted_zero_purchase,
        'Budget':m['budget']['amount'],'Budget scope':m['budget']['scope'],'Currency':m['budget']['currency'],'Purchase':package.cost_decomposition['purchase_vnd'],
        'Delivery':package.cost_decomposition['delivery_vnd'],'Committed cash':package.cost_decomposition['committed_procurement_vnd'],'Business ready':False,'Execution authorized':False}.items()])
    def records(name,headers,data):sheet(name,headers,[[r.get(k) for k in headers] for r in data])
    records('Orders',CUSTOMER_ORDER_FIELDS,rows)
    records('Delivery Schedule',['store_id','ingredient_id','ingredient_name','supplier_id','supplier_name','order_date','arrival_date','expiry_date','quantity','unit','pack_count','decision_stage'],sorted(rows,key=lambda r:(r['arrival_date'],r['ingredient_id'])))
    leaves=[]
    def flatten(a,path,value):
        if isinstance(value,dict) and value:
            for key,child in value.items():flatten(a,path+'.'+str(key),child)
        elif isinstance(value,list) and value:
            for i,child in enumerate(value):flatten(a,path+'.'+str(i),child)
        else:leaves.append({'assumption_id':a['assumption_id'],'scope':a.get('scope'),'field_path':path,'value':value})
    summaries=[]
    for a in conditions:
        value=a.get('value');flatten(a,a.get('field_path',''),value)
        summary=(f'{type(value).__name__}: {len(value)} entries; Assumption Values contains all leaf fields' if isinstance(value,(dict,list)) else value)
        summaries.append({**a,'value_summary':summary})
    records('Conditions',['assumption_id','kind','scope','field_path','value_summary','rationale','confirmation_required'],summaries)
    records('Assumption Values',['assumption_id','scope','field_path','value'],leaves)
    records('Validation',['rule_id','semantics','classification','target','unit','status','maximum_violation','weighted_violation'],package.rule_coverage)
    records('Daily Ledger',[f for f in LEDGER_FIELDS if f not in {'lots_consumed','lots_expired'}],ledger)
    traces=[]
    for row in ledger:
        for field in ['lots_consumed','lots_expired']:
            for trace in row.get(field,[]):
                traces.append({**trace,'date':row['date'],'ingredient_name':row['ingredient_name'],
                    'trace_type':'CONSUMPTION' if field=='lots_consumed' else 'EXPIRY',
                    'quantity':trace.get('quantity',trace.get('expired_quantity')),
                    'lot_expiry_date':trace.get('lot_expiry_date',trace.get('expiry_date')),
                    'probability_weight':row['probability_weight']})
    records('Lot Traces',['scenario_id','probability_weight','store_id','ingredient_id','ingredient_name','date','trace_type','lot_id','quantity','unit','lot_expiry_date'],traces)
    records('Terminal Expiry',['scenario_id','probability_weight','store_id','ingredient_id','ingredient_name','lot_id','unit','quantity_remaining','expiry_date','expiry_days_after_horizon','at_risk_terminal','certain_post_horizon_waste'],terminal['terminal_lots'])
    records('Diagnostics',['reason_code','field_paths','expected_meaning','action_required','details','violations','solver_status'],package.diagnostics)
    sheet('Worlds',['scenario_id','probability_weight'],[[k,v] for k,v in m['world_weights'].items()])
    if package.service_risk_metrics:records('Metrics',['store_id','ingredient_id','unit','expected_fill_rate','stockout_probability','expected_shortage','expected_expired_quantity','expected_ending_inventory'],package.service_risk_metrics['by_key'])
    wb.save(path);wb.close();check=load_workbook(path,data_only=False);errors=[];checked=0
    for name,data in expected.items():
        ws=check[name];observed=list(ws.values)
        if len(observed)!=len(data):errors.append(name+':row_count')
        if ws.freeze_panes!='A2' or not ws.auto_filter.ref:errors.append(name+':format')
        for actual,wanted in zip(observed,data,strict=True):
            for a,b in zip(actual,wanted,strict=True):
                if isinstance(b,(list,dict)):b=json.dumps(b,ensure_ascii=False,default=str)
                if isinstance(a,(date,datetime)):a=a.date().isoformat() if isinstance(a,datetime) else a.isoformat()
                if isinstance(b,float):
                    if a is None or abs(a-b)>max(1e-7,abs(b)*1e-12):errors.append(name+':numeric')
                elif a!=b and not (b=='' and a is None):errors.append(name+':value')
                checked+=1
        for row in ws:
            for cell in row:
                if cell.data_type in {'f','e'}:errors.append(name+':formula_or_error')
    check.close()
    if errors:raise ValueError('WORKBOOK_QA_FAILED:'+','.join(errors[:30]))
    return {'passed':True,'sheets':list(expected),'all_factual_cells_checked':checked,'order_rows':len(rows),
        'render_qa':'VALUE_AND_STYLE_READBACK_ONLY','limitation':'Excel/LibreOffice visual rendering unavailable; wrap/width/freeze/filter checked programmatically'}

def export_customer_plan(run,destination,accepted):
    destination=Path(destination);destination.mkdir(parents=True,exist_ok=True);folder=destination/'customer_plan'
    if folder.exists():raise FileExistsError('CUSTOMER_PACKAGE_EXISTS: use a fresh destination to preserve evidence')
    staging=destination/('.customer_plan_'+uuid.uuid4().hex);staging.mkdir()
    package,candidate=build_customer_view(run,accepted);terminal,ledger=perishable_evidence(run.request,candidate)
    write_json(staging/'customer_procurement_plan.json',package)
    csv_rows(staging/'customer_procurement_orders.csv',package.orders,CUSTOMER_ORDER_FIELDS)
    csv_rows(staging/'daily_inventory_ledger.csv',ledger,LEDGER_FIELDS)
    write_json(staging/'terminal_expiry.json',terminal)
    write_json(staging/'plan_conditions.json',{'business_ready':False,'execution_authorized':False,'conditions':package.assumptions_and_customer_conditions,
        'business_pending':run.request.business_issues,'supplier_terms':[o.model_dump(mode='json') for o in run.request.supplier_offers]})
    write_json(staging/'plan_validation.json',{'technical_outcome':package.technical_outcome,'critic':candidate.critic if candidate else None,
        'diagnostics':package.diagnostics,'request_hash':package.metadata['request_hash'],'exact_coverage':package.metadata['evaluation_coverage']})
    from shelfcash_forecast.optimization.parameter_registry import parameter_registry,recompute_objective
    write_json(staging/'parameter_registry.json',parameter_registry(run.request,run.config))
    if candidate:
        from shelfcash_forecast.optimization.strategies import default_strategy_profiles
        profiles={p.name:p for p in default_strategy_profiles()};profiles.update({p.name:p for p in run.request.strategy_profiles})
        write_json(staging/'independent_objective_check.json',recompute_objective(run.request,candidate,profiles[candidate.plan.strategy]))
    (staging/'CUSTOMER_PLAN_VI.md').write_text(explanation(package),encoding='utf-8')
    qa=write_workbook(staging/'customer_procurement_plan.xlsx',package,package.orders,package.assumptions_and_customer_conditions,ledger,terminal)
    write_json(staging/'workbook_verification.json',qa)
    os.replace(staging,folder)
    return package


def export_invalid_customer_diagnostics(result,destination):
    """An invalid payload has no bound horizon/terms; do not fabricate them."""
    from openpyxl import Workbook,load_workbook
    destination=Path(destination);folder=destination/'customer_plan'
    if folder.exists():raise FileExistsError('CUSTOMER_PACKAGE_EXISTS')
    staging=destination/('.customer_plan_'+uuid.uuid4().hex);staging.mkdir()
    diagnostics=[d.model_dump(mode='json') for d in result.diagnostics]
    package={'schema_version':2,'request_id':result.request_id,'binding_status':'UNBOUND_INVALID_REQUEST',
        'technical_outcome':result.technical_outcome,'orders':[],'accepted_zero_purchase':False,'diagnostics':diagnostics,
        'business_ready':False,'execution_authorized':False,'metadata':None,'service_risk_metrics':None}
    write_json(staging/'customer_procurement_plan.json',package)
    csv_rows(staging/'customer_procurement_orders.csv',[],CUSTOMER_ORDER_FIELDS)
    csv_rows(staging/'diagnostics.csv',diagnostics,['reason_code','field_paths','value','expected_meaning','action_required'])
    wb=Workbook();ws=wb.active;ws.title='Summary';ws.append(['Field','Value']);ws.append(['Status',result.technical_outcome]);ws.append(['Binding','UNBOUND_INVALID_REQUEST']);ws.append(['Accepted zero purchase',False])
    orders=wb.create_sheet('Orders');orders.append(CUSTOMER_ORDER_FIELDS)
    sheet=wb.create_sheet('Diagnostics');headers=['reason_code','field_paths','value','expected_meaning','action_required'];sheet.append(headers)
    for d in diagnostics:sheet.append([json.dumps(d.get(h),ensure_ascii=False,default=str) for h in headers])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for row in sheet:
            for c in row:
                if isinstance(c.value,str):c.data_type='s'
    wb.save(staging/'customer_procurement_plan.xlsx');wb.close()
    check=load_workbook(staging/'customer_procurement_plan.xlsx');assert list(check['Orders'].values)[0]==tuple(CUSTOMER_ORDER_FIELDS);check.close()
    text='INVALID_INPUT; chưa có accepted plan; không thể bind snapshot/horizon/terms từ payload chưa hợp lệ.\n\n'+'\n'.join(d.reason_code+': '+str(d.field_paths)+'; '+str(d.expected_meaning) for d in result.diagnostics)
    (staging/'CUSTOMER_PLAN_VI.md').write_text(text,encoding='utf-8');write_json(staging/'plan_validation.json',{'technical_outcome':result.technical_outcome,'diagnostics':diagnostics})
    os.replace(staging,folder)
    return {'status':'COMPLETE','customer_export_complete':True,'schema_version':2,'binding_status':'UNBOUND_INVALID_REQUEST'}

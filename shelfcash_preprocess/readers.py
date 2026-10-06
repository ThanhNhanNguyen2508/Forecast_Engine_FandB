from __future__ import annotations

import csv
import hashlib
import io
import json
import mimetypes
import shutil
import stat
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

import pandas as pd

from shelfcash_preprocess.config import Settings
from shelfcash_preprocess.models import SourceFile


SUPPORTED_SUFFIXES = {
    ".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".json", ".jsonl",
    ".parquet", ".pdf", ".png", ".jpg", ".jpeg", ".zip",
}


class ReaderError(RuntimeError):
    code = "READER_ERROR"


class ReaderUnavailable(ReaderError):
    code = "READER_UNAVAILABLE"


def _limit(settings: Settings | None, name: str, default: int) -> int:
    return int(getattr(settings, name, default))


def _check_table_shape(
    rows: int, columns: int, settings: Settings | None, *, source: str
) -> None:
    max_rows = _limit(settings, "max_table_rows", 1_000_000)
    max_columns = _limit(settings, "max_table_columns", 10_000)
    max_cells = _limit(settings, "max_table_cells", 10_000_000)
    if rows > max_rows:
        raise ReaderError(f"TABLE_ROW_LIMIT_EXCEEDED:{source}:{rows}>{max_rows}")
    if columns > max_columns:
        raise ReaderError(
            f"TABLE_COLUMN_LIMIT_EXCEEDED:{source}:{columns}>{max_columns}"
        )
    if rows and columns > max_cells // rows:
        raise ReaderError(
            f"TABLE_CELL_LIMIT_EXCEEDED:{source}:{rows}x{columns}>{max_cells}"
        )


@dataclass
class RawGrid:
    source_id: str
    source_path: str
    container: str
    container_kind: str
    rows: list[list[Any]]
    hidden: bool = False
    formulas: dict[str, str] = field(default_factory=dict)
    merged_ranges: list[str] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _source(path: Path, base: Path, archive_parent: str | None = None) -> SourceFile:
    digest = sha256_file(path)
    return SourceFile(
        source_id=f"src_{digest[:16]}",
        relative_path=path.relative_to(base).as_posix() if path.is_relative_to(base) else path.name,
        sha256=digest,
        size_bytes=path.stat().st_size,
        media_type=mimetypes.guess_type(path.name)[0] or "application/octet-stream",
        archive_parent=archive_parent,
    )


def inventory_input(
    input_path: str | Path,
    workspace: Path,
    settings: Settings,
) -> tuple[list[SourceFile], list[Path], list[str]]:
    """Create an immutable inventory and safely expand ZIPs into the run workspace."""

    source = Path(input_path).resolve()
    if not source.exists():
        raise FileNotFoundError(source)
    base = source if source.is_dir() else source.parent
    candidates = [source] if source.is_file() else sorted(
        p for p in source.rglob("*") if p.is_file()
    )
    if len(candidates) > settings.max_files:
        raise ReaderError(f"Too many input files: {len(candidates)} > {settings.max_files}")
    total_input_bytes = sum(path.stat().st_size for path in candidates)
    if total_input_bytes > settings.max_total_input_bytes:
        raise ReaderError(
            "TOTAL_INPUT_BYTES_LIMIT_EXCEEDED:"
            f"{total_input_bytes}>{settings.max_total_input_bytes}"
        )
    inventory: list[SourceFile] = []
    expanded: list[Path] = []
    warnings: list[str] = []
    archive_dir = workspace / "raw" / "archives"
    archive_dir.mkdir(parents=True, exist_ok=True)
    seen_hashes: set[str] = set()
    expanded_total = 0
    expanded_count = 0
    for path in candidates:
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            warnings.append(f"UNSUPPORTED_FILE:{path.name}")
            continue
        if path.stat().st_size > settings.max_file_bytes:
            raise ReaderError(f"File exceeds size limit: {path.name}")
        item = _source(path, base)
        if item.sha256 in seen_hashes:
            inventory.append(item.model_copy(update={"status": "DUPLICATE_SKIPPED", "notes": ["Identical bytes already inventoried in this run"]}))
            warnings.append(f"IDENTICAL_SOURCE_SKIPPED:{path.name}:{item.sha256[:16]}")
            continue
        seen_hashes.add(item.sha256)
        inventory.append(item)
        if suffix != ".zip":
            expanded.append(path)
            continue
        target = archive_dir / item.source_id
        target.mkdir(parents=True, exist_ok=True)
        total = 0
        count = 0
        with zipfile.ZipFile(path) as archive:
            for info in archive.infolist():
                if info.is_dir():
                    continue
                member = Path(info.filename.replace("\\", "/"))
                if member.is_absolute() or ".." in member.parts:
                    raise ReaderError(f"ARCHIVE_PATH_TRAVERSAL:{info.filename}")
                if member.suffix.lower() == ".zip":
                    raise ReaderError(f"NESTED_ARCHIVE_FORBIDDEN:{info.filename}")
                if info.flag_bits & 0x1:
                    raise ReaderError(f"ENCRYPTED_ARCHIVE_MEMBER_FORBIDDEN:{info.filename}")
                mode = (info.external_attr >> 16) & 0xF000
                if mode == stat.S_IFLNK:
                    raise ReaderError(f"ARCHIVE_SYMLINK_FORBIDDEN:{info.filename}")
                count += 1
                expanded_count += 1
                total += info.file_size
                expanded_total += info.file_size
                if info.file_size > settings.max_file_bytes:
                    raise ReaderError(f"ARCHIVE_MEMBER_FILE_LIMIT_EXCEEDED:{info.filename}")
                if (
                    expanded_count > settings.max_files
                    or expanded_total > settings.max_archive_bytes
                ):
                    raise ReaderError(
                        "ARCHIVE_EXPANSION_LIMIT_EXCEEDED_ACROSS_INPUT_SET"
                    )
                destination = (target / member).resolve()
                if target.resolve() not in destination.parents:
                    raise ReaderError(f"ARCHIVE_PATH_TRAVERSAL:{info.filename}")
                destination.parent.mkdir(parents=True, exist_ok=True)
                with archive.open(info) as incoming, destination.open("wb") as outgoing:
                    shutil.copyfileobj(incoming, outgoing)
                if destination.suffix.lower() in SUPPORTED_SUFFIXES - {".zip"}:
                    extracted_item = _source(destination, target, item.relative_path)
                    if extracted_item.sha256 in seen_hashes:
                        inventory.append(extracted_item.model_copy(update={"status": "DUPLICATE_SKIPPED", "notes": ["Identical bytes already inventoried in this run"]}))
                        warnings.append(f"IDENTICAL_SOURCE_SKIPPED:{info.filename}:{extracted_item.sha256[:16]}")
                    else:
                        seen_hashes.add(extracted_item.sha256)
                        expanded.append(destination)
                        inventory.append(extracted_item)
    return inventory, expanded, warnings


def _trim_rows(rows: list[list[Any]]) -> list[list[Any]]:
    width = max((len(row) for row in rows), default=0)
    padded = [row + [None] * (width - len(row)) for row in rows]
    while padded and all(value in (None, "") for value in padded[-1]):
        padded.pop()
    while width and all(row[width - 1] in (None, "") for row in padded):
        width -= 1
    return [row[:width] for row in padded]


def _read_excel(
    path: Path, source_id: str, settings: Settings | None = None
) -> list[RawGrid]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ReaderUnavailable(
            "XLSX_UNAVAILABLE: install with `pip install openpyxl`"
        ) from exc
    formulas_book = load_workbook(path, read_only=False, data_only=False, keep_vba=False)
    cached_book = load_workbook(path, read_only=False, data_only=True, keep_vba=False)
    grids: list[RawGrid] = []
    for sheet in formulas_book.worksheets:
        _check_table_shape(
            sheet.max_row, sheet.max_column, settings, source=f"{path.name}:{sheet.title}"
        )
        cached = cached_book[sheet.title]
        rows: list[list[Any]] = []
        formulas: dict[str, str] = {}
        for r in range(1, sheet.max_row + 1):
            values: list[Any] = []
            for c in range(1, sheet.max_column + 1):
                cell = sheet.cell(r, c)
                value = cell.value
                if cell.data_type == "f":
                    formulas[cell.coordinate] = str(value)
                    value = cached.cell(r, c).value
                values.append(value)
            rows.append(values)
        grids.append(
            RawGrid(
                source_id=source_id,
                source_path=str(path),
                container=sheet.title,
                container_kind="excel_sheet",
                rows=_trim_rows(rows),
                hidden=sheet.sheet_state != "visible",
                formulas=formulas,
                merged_ranges=[str(item) for item in sheet.merged_cells.ranges],
                metadata={
                    "sheet_state": sheet.sheet_state,
                    "excel_epoch": str(formulas_book.epoch),
                    "macros_executed": False,
                    "hidden_sheet_policy": "included_and_flagged",
                },
            )
        )
    formulas_book.close()
    cached_book.close()
    return grids


def _read_xls(path: Path, source_id: str, settings: Settings | None = None) -> list[RawGrid]:
    try:
        import xlrd
    except ImportError as exc:
        raise ReaderUnavailable(
            "XLS_UNAVAILABLE: install with `pip install xlrd`"
        ) from exc
    book = xlrd.open_workbook(path, on_demand=True)
    grids: list[RawGrid] = []
    for sheet in book.sheets():
        _check_table_shape(
            sheet.nrows, sheet.ncols, settings, source=f"{path.name}:{sheet.name}"
        )
        rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
        grids.append(
            RawGrid(source_id, str(path), sheet.name, "xls_sheet", _trim_rows(rows))
        )
    book.release_resources()
    return grids


def _decode_text(path: Path) -> tuple[str, str, bool]:
    raw = path.read_bytes()
    candidates: list[tuple[str, str]] = []
    try:
        from charset_normalizer import from_bytes

        match = from_bytes(raw).best()
        if match is not None:
            candidates.append((str(match.encoding), str(match)))
    except ImportError:
        pass
    for encoding in ("utf-8-sig", "utf-8", "cp1258", "cp1252"):
        try:
            candidates.append((encoding, raw.decode(encoding)))
        except UnicodeDecodeError:
            continue
    if not candidates:
        raise ReaderError(f"ENCODING_UNDETERMINED:{path.name}")
    encoding, text = candidates[0]
    alternatives = {candidate[1] for candidate in candidates[1:3]}
    ambiguous = any(value != text for value in alternatives) and encoding.lower() not in {
        "utf_8", "utf-8", "utf-8-sig"
    }
    return text, encoding, ambiguous


def _read_delimited(
    path: Path, source_id: str, settings: Settings | None = None
) -> list[RawGrid]:
    text, encoding, ambiguous = _decode_text(path)
    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    rows: list[list[str]] = []
    maximum_columns = 0
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        maximum_columns = max(maximum_columns, len(row))
        _check_table_shape(
            len(rows) + 1, maximum_columns, settings, source=path.name
        )
        rows.append(list(row))
    return [
        RawGrid(
            source_id, str(path), path.name, "delimited", _trim_rows(rows),
            metadata={
                "encoding": encoding,
                "encoding_ambiguous": ambiguous,
                "delimiter": delimiter,
            },
        )
    ]


def _flatten_record(value: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    output: dict[str, Any] = {}
    for key, item in value.items():
        name = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(item, dict):
            output.update(_flatten_record(item, name))
        elif isinstance(item, list):
            output[name] = json.dumps(item, ensure_ascii=False, separators=(",", ":"))
        else:
            output[name] = item
    return output


def _records_grid(records: list[dict[str, Any]]) -> list[list[Any]]:
    flattened = [_flatten_record(record) for record in records]
    columns = sorted({key for record in flattened for key in record})
    return [columns, *[[record.get(column) for column in columns] for record in flattened]]


def _read_json(path: Path, source_id: str, settings: Settings | None = None) -> list[RawGrid]:
    text, encoding, ambiguous = _decode_text(path)
    if path.suffix.lower() == ".jsonl":
        records = [json.loads(line) for line in text.splitlines() if line.strip()]
    else:
        payload = json.loads(text)
        records = payload if isinstance(payload, list) else [payload]
    if not all(isinstance(item, dict) for item in records):
        raise ReaderError("JSON_RECORDS_REQUIRED: top level must be object or list of objects")
    rows = _records_grid(records)
    _check_table_shape(len(rows), max((len(row) for row in rows), default=0), settings, source=path.name)
    return [RawGrid(source_id, str(path), path.name, "json_records", rows, metadata={"encoding": encoding, "encoding_ambiguous": ambiguous, "nested_lists": "preserved_as_json"})]


def _read_parquet(path: Path, source_id: str, settings: Settings | None = None) -> list[RawGrid]:
    try:
        import pyarrow.parquet as pq
        metadata = pq.ParquetFile(path).metadata
        _check_table_shape(
            metadata.num_rows,
            metadata.num_columns,
            settings,
            source=path.name,
        )
        frame = pd.read_parquet(path)
    except ImportError as exc:
        raise ReaderUnavailable(
            "PARQUET_UNAVAILABLE: install with `pip install pyarrow`"
        ) from exc
    rows = [frame.columns.tolist(), *frame.astype(object).where(frame.notna(), None).values.tolist()]
    return [RawGrid(source_id, str(path), path.name, "parquet", rows)]


def _read_pdf(path: Path, source_id: str, settings: Settings | None = None) -> list[RawGrid]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise ReaderUnavailable(
            "PDF_UNAVAILABLE: install with `pip install pdfplumber`"
        ) from exc
    grids: list[RawGrid] = []
    with pdfplumber.open(path) as pdf:
        max_pages = _limit(settings, "max_pdf_pages", 500)
        if len(pdf.pages) > max_pages:
            raise ReaderError(f"PDF_PAGE_LIMIT_EXCEEDED:{len(pdf.pages)}>{max_pages}")
        for page_number, page in enumerate(pdf.pages, 1):
            tables = page.extract_tables() or []
            for index, table in enumerate(tables, 1):
                grids.append(RawGrid(source_id, str(path), f"page-{page_number}-table-{index}", "pdf_table", _trim_rows(table), metadata={"page": page_number, "extraction": "pdfplumber"}))
            if not tables:
                text = page.extract_text() or ""
                lines = [line.strip() for line in text.splitlines() if line.strip()]
                if lines:
                    rows = [[part for part in line.split("  ") if part] for line in lines]
                    grids.append(RawGrid(source_id, str(path), f"page-{page_number}-text", "pdf_text", _trim_rows(rows), metadata={"page": page_number, "extraction": "pdfplumber_text"}))
            if not tables and not (page.extract_text() or "").strip():
                try:
                    import pytesseract
                    image = page.to_image(resolution=300).original
                    text = pytesseract.image_to_string(image, lang="vie+eng")
                    lines = [line.strip() for line in text.splitlines() if line.strip()]
                    rows = [[part for part in line.split("  ") if part] for line in lines]
                    grids.append(RawGrid(source_id, str(path), f"page-{page_number}-scan", "pdf_scan_ocr",
                                         _trim_rows(rows), metadata={"page": page_number, "status": "OCR_PROCESSED",
                                                                            "backend": "pytesseract", "languages": ["vie", "eng"]}))
                except (ImportError, OSError, RuntimeError) as exc:
                    grids.append(RawGrid(source_id, str(path), f"page-{page_number}-scan", "pdf_scan", [],
                                         metadata={"page": page_number, "status": "OCR_UNAVAILABLE",
                                                   "action": "Install with pip install -e .[ocr] and install Tesseract vie language data.",
                                                   "reason": type(exc).__name__}))
    return grids


def _read_image(path: Path, source_id: str, settings: Settings | None = None) -> list[RawGrid]:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise ReaderUnavailable(
            "OCR_UNAVAILABLE: install package extras with `pip install -e .[ocr]` "
            "and install Tesseract with Vietnamese language data."
        ) from exc
    try:
        image = Image.open(path)
        max_pixels = _limit(settings, "max_image_pixels", 50_000_000)
        if image.width and image.height > max_pixels // image.width:
            raise ReaderError(
                f"IMAGE_PIXEL_LIMIT_EXCEEDED:{image.width}x{image.height}>{max_pixels}"
            )
        data = pytesseract.image_to_data(
            image, lang="vie+eng", output_type=pytesseract.Output.DICT
        )
    except ReaderError:
        raise
    except Exception as exc:
        raise ReaderUnavailable(
            "OCR_UNAVAILABLE: Tesseract executable or vie language model is missing"
        ) from exc
    lines: dict[tuple[int, int, int], list[tuple[int, str, float]]] = {}
    for index, text in enumerate(data["text"]):
        cleaned = str(text).strip()
        if not cleaned:
            continue
        key = (data["block_num"][index], data["par_num"][index], data["line_num"][index])
        lines.setdefault(key, []).append((int(data["left"][index]), cleaned, float(data["conf"][index])))
    rows = [[text for _left, text, _conf in sorted(words)] for words in lines.values()]
    confidences = [conf for words in lines.values() for _left, _text, conf in words]
    return [RawGrid(source_id, str(path), path.name, "ocr_image", _trim_rows(rows), metadata={"backend": "pytesseract", "languages": ["vie", "eng"], "mean_confidence": sum(confidences) / len(confidences) if confidences else 0.0})]


def read_file(
    path: Path, source_id: str, settings: Settings | None = None
) -> list[RawGrid]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _read_excel(path, source_id, settings)
    if suffix == ".xls":
        return _read_xls(path, source_id, settings)
    if suffix in {".csv", ".tsv"}:
        return _read_delimited(path, source_id, settings)
    if suffix in {".json", ".jsonl"}:
        return _read_json(path, source_id, settings)
    if suffix == ".parquet":
        return _read_parquet(path, source_id, settings)
    if suffix == ".pdf":
        return _read_pdf(path, source_id, settings)
    if suffix in {".png", ".jpg", ".jpeg"}:
        return _read_image(path, source_id, settings)
    raise ReaderError(f"Unsupported format: {suffix}")


def read_all(
    paths: Iterable[Path],
    inventory: list[SourceFile],
    settings: Settings | None = None,
) -> tuple[list[RawGrid], list[str]]:
    by_hash = {item.sha256: item for item in inventory}
    grids: list[RawGrid] = []
    issues: list[str] = []
    for path in paths:
        digest = sha256_file(path)
        item = by_hash.get(digest)
        if item is None:
            issues.append(f"SOURCE_NOT_IN_INVENTORY:{path.name}")
            continue
        try:
            loaded = read_file(path, item.source_id, settings)
            for grid in loaded:
                _check_table_shape(
                    len(grid.rows),
                    max((len(row) for row in grid.rows), default=0),
                    settings,
                    source=f"{path.name}:{grid.container}",
                )
                grid.metadata.setdefault("source_sha256", item.sha256)
            grids.extend(loaded)
            for grid in loaded:
                if grid.metadata.get("status") == "OCR_UNAVAILABLE":
                    issues.append(f"OCR_UNAVAILABLE:{path.name}:{grid.container}:{grid.metadata.get('action')}")
        except ReaderUnavailable as exc:
            issues.append(str(exc))
        except Exception as exc:
            issues.append(f"READ_FAILED:{path.name}:{type(exc).__name__}:{exc}")
    return grids, issues
